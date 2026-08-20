"""One-off script: builds master-resume.docx from the user's resume content."""
from docx import Document
from docx.shared import Pt

doc = Document()

name = doc.add_paragraph()
run = name.add_run("Ebube Okeke")
run.bold = True
run.font.size = Pt(16)

doc.add_paragraph("Houston, TX | (346) 397-5026 | gokeke05@gmail.com | linkedin.com/in/ebube-okeke-11ab94379")


def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)


def subheading(left, right):
    p = doc.add_paragraph()
    r = p.add_run(left)
    r.bold = True
    if right:
        p.add_run(f"\t{right}")


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


heading("EDUCATION")
subheading("University of Houston — Bachelor of Science, Biotechnology", "May 2028")
bullet(
    "Coursework: Intro to Biotechnology, Biopython, Microbiology & Lab, "
    "Biochemistry, Organic Chemistry, Organic Chemistry Lab"
)

heading("EXPERIENCE")
subheading("Immatics | CMC-GMP Manufacturing Assistant", "Jun 2025 - Jul 2025")
bullet(
    "Managed multiple time-sensitive tasks, including manufacturing "
    "documentation, patient records, and data entry, using Microsoft Excel "
    "and Word to maintain accurate records and support on-schedule cGMP "
    "production."
)
bullet(
    "Collaborated with manufacturing personnel by shadowing cleanroom "
    "production processes and reviewing batch records and Standard "
    "Operating Procedures (SOPs), building familiarity with cGMP "
    "compliance, manufacturing environments, and regulatory documentation "
    "practices."
)
bullet(
    "Supported accurate, audit-ready manufacturing operations by "
    "maintaining organized documentation and applying cGMP practices in a "
    "detail-oriented production environment."
)
bullet(
    "Partnered with the Manufacturing Science and Technology (MSAT) team "
    "on cross-functional process optimization initiatives, reviewing "
    "project information and supporting efforts to improve manufacturing "
    "robustness and regulatory compliance."
)

heading("PROJECTS")
subheading("Synthetic Bone Graft Scaffold", "Dec 2025 - Present")
bullet(
    "Designed a parametric family of synthetic bone graft scaffolds, "
    "including block, wedge, and cylinder geometries, in SolidWorks using "
    "size-configurable Design Tables and a porosity-representative feature "
    "grounded in real material standards."
)
bullet(
    "Developed a structured validation plan and dimensioned drawing "
    "package with GD&T tolerancing for cylindricity, position, and "
    "profile; FEA and physical testing are in progress."
)
bullet(
    "Created a standards-informed, configurable design foundation to "
    "support consistent future prototyping, validation, and adaptation "
    "across scaffold geometries."
)

subheading("Reinforcement Learning Control Agent", "Jan 2026 - May 2026")
bullet(
    "Developing a reinforcement learning control agent using PPO and DQN "
    "in Python within a Gym-based simulation environment, along with a "
    "training and evaluation pipeline to benchmark policy performance "
    "through iterative tuning."
)
bullet(
    "Built a reusable experimentation workflow that supports systematic "
    "comparison of control policies and informed iterative model "
    "improvements."
)

subheading("Pulse Oximeter (SpO2) Monitoring System", "Mar 2026 - Present")
bullet(
    "Designing a pulse oximeter (SpO2) monitoring system integrating a "
    "biomedical optics sensor with Arduino-based signal acquisition and a "
    "signal-processing pipeline to extract and validate SpO2 readings (in "
    "progress)."
)
bullet(
    "Established an integrated hardware-and-software foundation for "
    "validating reliable SpO2 monitoring workflows as development "
    "progresses."
)

heading("LEADERSHIP & COMMUNITY")
subheading(
    "National Society of Black Engineers (NSBE), University of Houston — "
    "Career Fair & Black in Tech Committee Member",
    "Sep 2024 - Present",
)
bullet(
    "Facilitated technical workshops on Python, HTML/CSS, and JavaScript "
    "for 40+ underrepresented students; supported outreach efforts that "
    "contributed to a 70% increase in chapter attendance."
)
bullet(
    "Co-organized UH's annual career fair connecting 30+ companies with "
    "200+ students; managed recruiter follow-up and tracked engagements "
    "in Excel, contributing to a 25% increase in student interview offers "
    "versus the prior year."
)
bullet(
    "Strengthened access to technical learning and career opportunities "
    "by supporting inclusive programming, employer engagement, and "
    "student follow-through."
)

subheading("Rosenberg-Richmond Helping Hand — Team Leader", "Jun 2023 - Aug 2025")
bullet(
    "Rebuilt a 15-year-old Google Calendar-based intake system as a food "
    "pantry allocation and appointment system resource tracker using "
    "FastAPI, PostgreSQL, Next.js, and OR-Tools CP-SAT, cutting fulfillment "
    "time to under 1 hour per submission while maintaining 100% "
    "documentation accuracy across 15+ families."
)
bullet(
    "Improved the pantry's ability to coordinate appointments and "
    "allocations by translating a manual scheduling process into a more "
    "structured digital workflow."
)

heading("SKILLS")
bullet("Programming: Python, Java, C++, C#, SQL, HTML/CSS, Biopython, MATLAB")
bullet("CAD & Engineering: SolidWorks, Fusion 360")
bullet(
    "Analysis & Quality: Statistical Analysis, Statistical Design Of "
    "Experiments (DOE), Manufacturing Practices (GMP), Clean Room "
    "Manufacturing, Manufacturing Environments"
)
bullet("Tools: Microsoft Office (Excel, Word, PowerPoint), Git, VS Code, Codex, Minitab")

heading("CERTIFICATIONS")
bullet("CITI Program (2026): Biomedical Research Ethics, Biosafety, Research Security")
bullet("MathWorks (2026): MATLAB Onramp & Simulink Onramp")
bullet("LinkedIn Learning: AutoCAD Essential Training, Fusion 360 Essential Training")

doc.save("master-resume.docx")
print("Saved master-resume.docx")
