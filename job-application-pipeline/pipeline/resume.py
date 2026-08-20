"""Tailors the master resume to a job description via the Anthropic API.

Only rephrases and reorders bullets that already exist in the master resume.
Never invents metrics, skills, or experience.
"""
import json
from dataclasses import dataclass
from pathlib import Path

from anthropic import Anthropic
from docx import Document

_BULLET_STYLE_HINTS = ("list", "bullet")

_SYSTEM_PROMPT = """You tailor resumes and draft outreach material. You are given the \
bullets from a candidate's master resume, grouped by section, and a target job \
description.

Rules (do not break these):
1. Never invent, exaggerate, or add metrics, skills, tools, or experience that \
are not already present in the given bullets.
2. You may rephrase a bullet's wording for clarity and impact, and reorder \
bullets within a section to foreground the most relevant ones.
3. Where natural, echo the job description's own action verbs and terminology \
(light-to-medium keyword overlap for ATS), but only when it still accurately \
describes what the bullet already says.
4. Do not merge, split, or delete bullets. Every input bullet must appear \
exactly once in the output, for the same section.
5. Separately, pick up to 3 of the ORIGINAL (unrevised) input bullets that \
contain a concrete number or metric and are most relevant to the job \
description. Quote them verbatim, character-for-character as given — do not \
edit them here. If fewer than 3 input bullets contain a number, return fewer.
6. Write a 2-3 sentence "fit paragraph" connecting the candidate's existing, \
given experience to what the job description asks for. It must only \
reference things actually present in the given bullets — no fabrication.

Return ONLY JSON of the form:
{"sections": [{"section": "<name>", "bullets": [{"original_index": <int>, "text": "<revised text>"}, ...]}],
 "email_highlights": ["<verbatim bullet 1>", ...],
 "fit_paragraph": "<2-3 sentences>"}
"""


def _iter_paragraphs_with_sections(doc: Document):
    """Groups bullets by contiguous run under the same subheading (e.g. one
    job or one project), never across a subheading/date-line boundary — a
    resume section like PROJECTS or EXPERIENCE holds several such runs, and
    reordering across them would scramble bullets out from under the wrong
    job/project.
    """
    current_section = "General"
    current_subheading = ""
    group_key = None
    in_bullet_run = False
    for index, paragraph in enumerate(doc.paragraphs):
        style_name = (paragraph.style.name or "").lower()
        text = paragraph.text.strip()
        if not text:
            continue
        is_heading = "heading" in style_name or (
            text.isupper() and len(text.split()) <= 6
        )
        is_bullet = any(hint in style_name for hint in _BULLET_STYLE_HINTS)
        if is_heading:
            current_section = text
            current_subheading = ""
            in_bullet_run = False
            continue
        if is_bullet:
            if not in_bullet_run:
                group_key = f"{current_section} :: {current_subheading or index}"
                in_bullet_run = True
            yield index, group_key, text
        else:
            current_subheading = text
            in_bullet_run = False


def _extract_bullets(doc: Document) -> dict[str, list[tuple[int, str]]]:
    sections: dict[str, list[tuple[int, str]]] = {}
    for index, section, text in _iter_paragraphs_with_sections(doc):
        sections.setdefault(section, []).append((index, text))
    return sections


def _call_claude(client: Anthropic, sections: dict, job_description: str, company: str) -> dict:
    sections_payload = [
        {"section": name, "bullets": [{"original_index": i, "text": t} for i, t in bullets]}
        for name, bullets in sections.items()
    ]
    user_message = (
        f"Target company: {company}\n\n"
        f"Job description:\n{job_description[:8000]}\n\n"
        f"Master resume bullets by section:\n{json.dumps(sections_payload, indent=2)}"
    )
    response = client.messages.create(
        model="claude-sonnet-4-5",
        max_tokens=4096,
        system=_SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user_message}],
    )
    raw_text = "".join(block.text for block in response.content if block.type == "text")
    raw_text = raw_text.strip()
    if raw_text.startswith("```"):
        raw_text = raw_text.split("```")[1]
        if raw_text.startswith("json"):
            raw_text = raw_text[4:]
    return json.loads(raw_text)


def _apply_revisions(doc: Document, revisions: dict) -> Document:
    paragraphs_by_index = {i: p for i, p in enumerate(doc.paragraphs)}

    for section in revisions.get("sections", []):
        ordered_indices = [b["original_index"] for b in section["bullets"]]
        text_by_index = {b["original_index"]: b["text"] for b in section["bullets"]}

        # Rewrite text in place (preserve the first run's formatting).
        for index in ordered_indices:
            paragraph = paragraphs_by_index.get(index)
            if paragraph is None or not paragraph.runs:
                continue
            paragraph.runs[0].text = text_by_index[index]
            for extra_run in paragraph.runs[1:]:
                extra_run.text = ""

        # Reorder paragraph XML elements within the section to match the
        # model's requested order, if it differs from the original order.
        original_order = sorted(ordered_indices)
        if ordered_indices == original_order:
            continue
        elements = [paragraphs_by_index[i]._p for i in original_order]
        anchor = elements[0]
        parent = anchor.getparent()
        insert_pos = list(parent).index(anchor)
        for element in elements:
            parent.remove(element)
        for offset, index in enumerate(ordered_indices):
            parent.insert(insert_pos + offset, paragraphs_by_index[index]._p)

    return doc


@dataclass
class TailoredResume:
    docx_path: Path
    email_highlights: list[str]
    fit_paragraph: str


def tailor_resume(master_resume_path: Path, job_description: str, company: str, api_key: str, output_path: Path) -> TailoredResume:
    client = Anthropic(api_key=api_key)
    doc = Document(str(master_resume_path))

    sections = _extract_bullets(doc)
    if not sections:
        raise ValueError(
            "No bulleted paragraphs found in the master resume. Make sure bullet "
            "points use a Word 'List Bullet'-style paragraph style."
        )

    revisions = _call_claude(client, sections, job_description, company)
    tailored_doc = _apply_revisions(doc, revisions)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    tailored_doc.save(str(output_path))

    return TailoredResume(
        docx_path=output_path,
        email_highlights=revisions.get("email_highlights", []),
        fit_paragraph=revisions.get("fit_paragraph", ""),
    )
